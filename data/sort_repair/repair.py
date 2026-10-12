# SPDX-FileCopyrightText: 2026 Cade Summers
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Undo the re-sort that scrambled Multi-stars AQ:AT and stranded both star tabs' provenance borders.

    python data/sort_repair/repair.py          plan: print every change; touches nothing
    python data/sort_repair/repair.py write    make them, note them on Sources & Provenance, rename

Between the 6c63081 snapshot (hy-brasil-20261002-2218) and the 0f2732b one (hy-brasil-20261004-1158), Multi-stars
and Systems were re-sorted through their basic filters. Two things went wrong:

- Multi-stars' filter ended at AP (endColumnIndex 42) on a 46-column tab, so the sort moved A:AP and left AQ:AT
  (Separation from / (AU) / type / source) where they were: 94 of 142 rows carried another star's separation, with
  that star's notes and borders. AQ:AT is a pure permutation of the old values, so the old row order is the map back.
- Borders never travel with a sort (notes do), so the orange and pink provenance boxes in A:AP on Multi-stars and in
  every column on Systems stayed on the row numbers, not the stars.

The truth for both is the 6c63081 snapshot, by star name. The repair:
1. AQ:AT back to their stars -- values, notes and borders in one atomic batch, through a scratch tab.
2. Multi-stars' basic filter widened to all 46 columns (same sort specs), so the next sort carries AQ:AT along.
3. Every orange/pink box, both tabs, put on the star it belongs to; stray boxes go back to the tab's plain rule.
4. Everything checked again by star name, then Sources & Provenance, then the rename.

Rerunnable: step 3 is computed from the live state each time, so after any later sort it re-homes the borders.
"""
import io
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "data" / "four_plus_stars"))
from hybrasil.data import google_sheets as gs  # noqa: E402
from fill import letter, q  # noqa: E402  (importing fill also installs its 429 retry wrapper)

PLAN = "write" not in sys.argv[1:]
SID = gs.resolve_spreadsheet_id()
API = gs._sheets_api()
MS, SY, SP = "Multi-stars within 100 ly", "Systems within 100 ly", "Sources & Provenance"
GOOD_REV, GOOD_FILE = "6c63081", "data/hy-brasil-20261002-2218.xlsx"
KEY = {MS: "Star (component)", SY: "System / Host"}
SEP0, SEP1 = 42, 46                       # AQ:AT, 0-based half-open
MARKS = {"FF9900", "FF3399"}              # orange inferred, pink Wikipedia
TMP_ID, TMP_TITLE = 987654321, "_sort_repair_tmp"
LOG = HERE / "repair_log.json"
DATE = "4 October 2026"


def a1_to_rc(a1):
    m = re.match(r"([A-Z]+)(\d+)$", a1)
    col = 0
    for ch in m[1]:
        col = col * 26 + ord(ch) - 64
    return int(m[2]), col - 1                       # 1-based row, 0-based column


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# ── the truth: the 6c63081 snapshot ─────────────────────────────────────────────────────────────────────────────
def good_tables():
    raw = subprocess.run(["git", "show", f"{GOOD_REV}:{GOOD_FILE}"], cwd=ROOT, capture_output=True, check=True).stdout
    wb_v = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    wb_f = openpyxl.load_workbook(io.BytesIO(raw))           # formats
    out = {}
    for tab in (MS, SY):
        rows = [list(r) for r in wb_v[tab].iter_rows(values_only=True)]
        head = rows[0]
        k = head.index(KEY[tab])
        data = [r for r in rows[1:] if r[0]]
        names = [r[k] for r in data]
        ws = wb_f[tab]
        marks = {}
        for i, name in enumerate(names, start=2):
            for c in range(len(head)):
                cell = ws.cell(i, c + 1)
                cols = set()
                for side in ("left", "right", "top", "bottom"):
                    sd = getattr(cell.border, side)
                    rgb = sd.color.rgb if sd is not None and sd.color is not None else None
                    cols.add(rgb[2:] if isinstance(rgb, str) else None)
                if len(cols) == 1 and cols <= MARKS:          # a box: all four sides the same marking colour
                    marks[(name, c)] = cols.pop()
        out[tab] = dict(head=head, names=names, rows=data, marks=marks)
    return out


# ── the live sheet ──────────────────────────────────────────────────────────────────────────────────────────────
def live_values(tab, ncols):
    v = API.values().get(spreadsheetId=SID, range=f"{q(tab)}!A1:{letter(ncols - 1)}3000",
                         valueRenderOption="UNFORMATTED_VALUE").execute()["values"]
    head = v[0]
    data = [r + [""] * (ncols - len(r)) for r in v[1:] if r and r[0] != ""]
    return head, data


def live_marks(tab, names, ncols):
    """{(name, col): colour} for every box, and the tab's plain rule (the commonest unmarked side)."""
    b = gs.get_borders(None, f"{q(tab)}!A2:{letter(ncols - 1)}{len(names) + 1}")
    marks, plain = {}, Counter()
    for a1, sides in b.items():
        r, c = a1_to_rc(a1)
        cols = {col for col, _ in sides.values()}
        if len(sides) == 4 and len(cols) == 1 and cols <= MARKS:
            marks[(names[r - 2], c)] = cols.pop()
        else:
            plain.update(s for s in sides.values() if s[0] not in MARKS)
    return marks, plain.most_common(1)[0][0]


def live_notes(tab, names, ncols):
    out = {}
    for a1, text in gs.get_notes(None, f"{q(tab)}!A2:{letter(ncols - 1)}{len(names) + 1}").items():
        r, c = a1_to_rc(a1)
        out[(r - 2, c)] = text
    return out


def border_fixes(tab, names, want, have):
    """(cells to re-rule plain, {cell: colour} to box) that turn `have` into `want`, both keyed by (name, col)."""
    row = {n: i + 2 for i, n in enumerate(names)}
    plain = sorted(f"{q(tab)}!{letter(c)}{row[n]}" for (n, c), col in have.items() if want.get((n, c)) != col)
    box = {f"{q(tab)}!{letter(c)}{row[n]}": col for (n, c), col in want.items() if have.get((n, c)) != col}
    return plain, box


def apply_borders(tab, names, plain, want, rule):
    """Strays to the plain rule, then EVERY wanted box, not only the changed ones: a box shares its edges with
    its neighbours, so re-ruling (or pasting) the cell next door strips a side off a box that was right."""
    colour, style = rule
    if plain:
        gs.set_border(SID, plain, colour or gs.PLAIN_RULE, style=style)
    row = {n: i + 2 for i, n in enumerate(names)}
    for colour in sorted(set(want.values())):
        gs.set_border(SID, [f"{q(tab)}!{letter(c)}{row[n]}" for (n, c), col in want.items() if col == colour],
                      "FF" + colour)


def main():
    good = good_tables()
    props = {s["properties"]["title"]: s for s in
             API.get(spreadsheetId=SID, fields="sheets(properties(title,sheetId,gridProperties),basicFilter)"
                     ).execute()["sheets"]}
    assert TMP_TITLE not in props, "a scratch tab from an earlier run is still there"
    state = {}
    for tab in (MS, SY):
        g = good[tab]
        ncols = len(g["head"])
        head, data = live_values(tab, ncols)
        assert head == g["head"], f"{tab}: headers changed since {GOOD_REV}"
        names = [r[head.index(KEY[tab])] for r in data]
        assert len(set(names)) == len(names), f"{tab}: duplicate names"
        assert set(names) == set(g["names"]), (tab, set(names) ^ set(g["names"]))
        state[tab] = dict(head=head, data=data, names=names, ncols=ncols)

    # 1. Multi-stars AQ:AT: live row i still holds what the snapshot's row i holds -- that is the map back.
    # Skipped once AQ:AT already sits with its stars (a rerun after the filter fix).
    ms, gms = state[MS], good[MS]
    live_row = {n: i + 2 for i, n in enumerate(ms["names"])}
    snap_sep = {n: [norm(x) for x in r[SEP0:SEP1]] for n, r in zip(gms["names"], gms["rows"])}
    aligned = all([norm(x) for x in r[SEP0:SEP1]] == snap_sep[n] for n, r in zip(ms["names"], ms["data"]))
    if aligned:
        moves = []
    else:
        for i, (lr, gr) in enumerate(zip(ms["data"], gms["rows"])):
            assert [norm(x) for x in lr[SEP0:SEP1]] == [norm(x) for x in gr[SEP0:SEP1]], (
                f"AQ:AT row {i + 2} is not the snapshot's row {i + 2}: not a pure permutation, stop")
        moves = [(gi + 2, live_row[n], n) for gi, n in enumerate(gms["names"]) if gi + 2 != live_row[n]]
    pre_notes = live_notes(MS, ms["names"], ms["ncols"])
    # Where every note should end up, by name: A:AP notes travelled with their rows, AQ:AT ones belong to the
    # star the snapshot had on that row.
    want_notes = {}
    for (r, c), text in pre_notes.items():
        want_notes[(gms["names"][r] if c >= SEP0 and moves else ms["names"][r], c)] = text

    flt = props[MS]["basicFilter"]
    assert flt["range"]["endColumnIndex"] in (SEP0, SEP1), flt["range"]
    new_flt = dict(flt, range=dict(flt["range"], endColumnIndex=SEP1))

    # 3 (planned). Borders by name. AQ:AT's are carried by step 1's paste, so plan them as already moved.
    fixes = {}
    for tab in (MS, SY):
        st = state[tab]
        have, rule = live_marks(tab, st["names"], st["ncols"])
        if tab == MS and moves:
            have = {((gms["names"][live_row[n] - 2] if c >= SEP0 else n), c): col for (n, c), col in have.items()}
        fixes[tab] = (*border_fixes(tab, st["names"], good[tab]["marks"], have), rule, have)

    bad_now = sum(1 for r in ms["data"] if r[SEP0] and r[SEP0] not in
                  {x[1] for x in ms["data"] if x[0] == r[0]})
    print(f"Multi-stars: {len(moves)} of {len(ms['names'])} rows get their AQ:AT back; "
          f"{bad_now} 'Separation from' values name a star outside their system now")
    for g_row, l_row, n in moves[:8]:
        print(f"   {n:28} AQ:AT from row {g_row} -> row {l_row}")
    print("   ...")
    print(f"   notes in AQ:AT moving: {sum(1 for (r, c) in pre_notes if c >= SEP0 and gms['names'][r] != ms['names'][r])}")
    print(f"Multi-stars filter: columns A:{letter(SEP0 - 1)} -> A:{letter(SEP1 - 1)}, "
          f"sort specs kept: {flt.get('sortSpecs')}")
    for tab in (MS, SY):
        plain, box, rule, have = fixes[tab]
        print(f"{tab}: {len(good[tab]['marks'])} boxes in the snapshot, {len(have)} live; "
              f"{len(plain)} stray boxes back to the plain rule {rule}, {len(box)} boxes restored")
        for a in sorted(box)[:6]:
            print("   box", a, box[a])
        for a in plain[:6]:
            print("   plain", a)
    differ = sum(1 for g_row, l_row, _ in moves if [norm(x) for x in gms["rows"][g_row - 2][SEP0:SEP1]] !=
                 [norm(x) for x in gms["rows"][l_row - 2][SEP0:SEP1]])
    lines = sp_lines(len(moves), differ, fixes, good)
    print("\n".join(lines))
    if PLAN:
        return

    # ── writes ──────────────────────────────────────────────────────────────────────────────────────────────
    if not moves and flt["range"]["endColumnIndex"] == SEP1 and not any(f[0] or f[1] for f in fixes.values()):
        print("nothing to repair")
        return
    # Guard, as update_cells' expect would: nothing on either tab has moved since the reads above.
    for tab in (MS, SY):
        assert live_values(tab, state[tab]["ncols"])[1] == state[tab]["data"], f"{tab} changed since the read; rerun"

    # 1 and 2.
    if moves:
        move_back(props[MS]["properties"]["sheetId"], len(ms["names"]), moves)
    if flt["range"]["endColumnIndex"] != SEP1:
        API.batchUpdate(spreadsheetId=SID, body={"requests": [{"setBasicFilter": {"filter": new_flt}}]}).execute()
    gs.invalidate_cache(SID)

    # 3. Borders, from the live state now (the filter call may have re-sorted), by name.
    done = {}
    for tab in (MS, SY):
        head, data = live_values(tab, state[tab]["ncols"])
        names = [r[head.index(KEY[tab])] for r in data]
        have, rule = live_marks(tab, names, state[tab]["ncols"])
        plain, box = border_fixes(tab, names, good[tab]["marks"], have)
        apply_borders(tab, names, plain, good[tab]["marks"], rule)
        done[tab] = (len(plain), len(box))

    verify(state, good, flt, want_notes)

    # Sources & Provenance: grow the grid, then explicit ranges (never values().append()).
    meta = API.get(spreadsheetId=SID, ranges=[f"{q(SP)}!A1"], fields="sheets(properties(sheetId,gridProperties))"
                   ).execute()["sheets"][0]["properties"]
    end = meta["gridProperties"]["rowCount"]
    assert gs.read_cells(SID, [f"{q(SP)}!A{end}"])[f"{q(SP)}!A{end}"][0][0].strip()
    sp0, sp1 = end + 1, end + len(lines)
    API.batchUpdate(spreadsheetId=SID, body={"requests": [{"appendDimension": {
        "sheetId": meta["sheetId"], "dimension": "ROWS", "length": len(lines)}}]}).execute()
    gs.update_cells(SID, {f"{q(SP)}!A{sp0}:A{sp1}": [[x] for x in lines]},
                    expect={f"{q(SP)}!A{i}": "" for i in (sp0, sp0 + 1, sp1)})

    title = gs.rename_workbook(None, gs.versioned_title())
    gs.invalidate_cache(SID)
    log = {"aq_at_rows_restored": len(moves), "filter_columns": [0, SEP1], "sp_rows": [sp0, sp1], "title": title,
           "borders": {t: {"re_ruled_plain": p, "boxed": b} for t, (p, b) in done.items()}}
    LOG.write_text(json.dumps(log, indent=1), encoding="utf-8")
    print("done:", json.dumps(log))


def move_back(mid, n, moves):
    """AQ:AT row by row from where the snapshot had it to where its star is now: values, notes and formats.

    Copied out to a hidden scratch tab first, since every source row is also some other row's destination; one
    batchUpdate, so it lands whole or not at all."""
    def grid(sheet, r0, r1, c0, c1):
        return {"sheetId": sheet, "startRowIndex": r0, "endRowIndex": r1, "startColumnIndex": c0,
                "endColumnIndex": c1}

    w = SEP1 - SEP0
    reqs = [{"addSheet": {"properties": {"sheetId": TMP_ID, "title": TMP_TITLE, "hidden": True,
                                         "gridProperties": {"rowCount": n + 2, "columnCount": w}}}},
            {"copyPaste": {"source": grid(mid, 1, n + 1, SEP0, SEP1), "destination": grid(TMP_ID, 1, n + 1, 0, w),
                           "pasteType": "PASTE_NORMAL"}}]
    for g_row, l_row, _ in moves:
        reqs.append({"copyPaste": {"source": grid(TMP_ID, g_row - 1, g_row, 0, w),
                                   "destination": grid(mid, l_row - 1, l_row, SEP0, SEP1),
                                   "pasteType": "PASTE_NORMAL"}})
    reqs.append({"deleteSheet": {"sheetId": TMP_ID}})
    API.batchUpdate(spreadsheetId=SID, body={"requests": reqs}).execute()


def verify(state, good, flt, want_notes):
    """Every value, note and box as it should be, by name; every 'Separation from' inside its own system."""
    for tab in (MS, SY):
        head, data = live_values(tab, state[tab]["ncols"])
        k = head.index(KEY[tab])
        before = {r[k]: [norm(x) for x in r] for r in state[tab]["data"]}
        snap = {r[k]: [norm(x) for x in r] for r in good[tab]["rows"]}
        for r in data:
            now = [norm(x) for x in r]
            if tab == MS:
                assert now[:SEP0] == before[r[k]][:SEP0], f"{r[k]}: A:AP changed"
                assert now[SEP0:SEP1] == snap[r[k]][SEP0:SEP1], f"{r[k]}: AQ:AT not restored"
            else:
                assert now == before[r[k]], f"{r[k]}: Systems row changed"
        names = [r[k] for r in data]
        have, _ = live_marks(tab, names, state[tab]["ncols"])
        assert have == good[tab]["marks"], (tab, set(have.items()) ^ set(good[tab]["marks"].items()))
        if tab == MS:
            got = {(names[r], c): t for (r, c), t in live_notes(MS, names, state[MS]["ncols"]).items()}
            assert got == want_notes, set(got.items()) ^ set(want_notes.items())
            members = {}
            for r in data:
                members.setdefault(r[0], set()).add(r[k])
            stray = [r[k] for r in data if r[SEP0] and r[SEP0] not in members[r[0]]]
            assert not stray, stray
    f_now = [s for s in API.get(spreadsheetId=SID, fields="sheets(properties(title),basicFilter)").execute()["sheets"]
             if s["properties"]["title"] == MS][0]["basicFilter"]
    assert f_now["range"]["endColumnIndex"] == SEP1 and f_now.get("sortSpecs") == flt.get("sortSpecs"), f_now


def sp_lines(n_moves, differ, fixes, good):
    ms_plain, ms_box = len(fixes[MS][0]), len(fixes[MS][1])
    sy_plain, sy_box = len(fixes[SY][0]), len(fixes[SY][1])
    return [
        "",
        f"MULTI-STARS AND SYSTEMS UN-SCRAMBLED AFTER A RE-SORT, {DATE.upper()}",
        "  WHAT WENT WRONG: between the hy-brasil-20261002-2218 and hy-brasil-20261004-1158 exports both tabs were "
        "re-sorted through their basic filters. Multi-stars' filter covered A:AP of A:AT, so the sort left "
        f"'Separation from', 'Separation (AU)', 'Separation type' and 'Separation source' behind: {n_moves} of "
        f"{len(good[MS]['names'])} rows were parted from those cells, {differ} of them showing another star's "
        "separation (e.g. GJ 896 A '127.84 AU from HD 149806'), "
        "with that star's notes and borders. The viewer did not complain: a companion whose 'Separation from' is "
        "not in its own system is drawn round the primary. Borders never travel with a sort, so the orange and "
        "pink provenance boxes on both tabs had also stayed on their row numbers.",
        "  FIX (data/sort_repair/repair.py, by Claude): the hy-brasil-20261002-2218 export is the reference, by star "
        "name. Multi-stars AQ:AT was a pure permutation of it, so each row's four cells, notes and borders were "
        "copied back to their star in one batch. The filter now covers A:AT with the same sort, so a future sort "
        f"carries them along. Provenance boxes re-homed by name: Multi-stars {ms_box} restored, {ms_plain} strays "
        f"re-ruled plain; Systems {sy_box} restored, {sy_plain} strays re-ruled plain. Checked afterwards by name: "
        "every value, note and box as the reference has it, and every 'Separation from' names a star of its own "
        "system. No value was changed.",
        "  AFTER ANY RE-SORT: rerun repair.py; its border step is computed from the live sheet and re-homes the "
        "boxes again. The reference export must then be one taken after the last intended border change.",
    ]


if __name__ == "__main__":
    main()
