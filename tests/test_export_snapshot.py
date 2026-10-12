# SPDX-FileCopyrightText: 2026 Cade Summers
# SPDX-License-Identifier: AGPL-3.0-or-later
"""The snapshot exporter (hybrasil/data/export_snapshot.py), with no network.

Every workbook here is built in the test and served by a fake `fetch`, so
what is pinned is the decision: a new snapshot only when content changed,
the old one removed, and the README row naming the new one.
"""

import datetime as dt
import io
import json
import time

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Border, Side

from hybrasil.data import export_snapshot as ex

ORANGE = "FFFF9900"
OLD = "hybrasil-20261010-2036"
NEW = "hybrasil-20261011-1836"
TODAY = dt.date(2026, 10, 12)

README = f"""# Workbook snapshots

| File | Workbook name when exported | Drive file ID | Exported | How |
| --- | --- | --- | --- | --- |
| `{OLD}.xlsx` | `{OLD}` | private (`sheet_id.txt`) | 2026-10-10 | by hand |

The middle column is the name the workbook carried at export time.
"""


def _xlsx(value="Sun", border=None, note=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Systems"
    ws["A1"], ws["B1"] = "Name", "Luminosity"
    ws["A2"], ws["B2"] = value, 1.0
    if border:
        side = Side(style="thin", color=border)
        ws["B2"].border = Border(left=side, right=side, top=side, bottom=side)
    if note:
        ws["B2"].comment = Comment(note, "hybrasil")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def data_dir(tmp_path):
    (tmp_path / f"{OLD}.xlsx").write_bytes(_xlsx())
    (tmp_path / "README.md").write_text(README, encoding="utf-8")
    return tmp_path


def _fetch(title, blob):
    def fetch(url, token=None):
        if "sheets.googleapis.com" in url:
            return json.dumps({"properties": {"title": title}}).encode()
        assert "export?format=xlsx" in url
        return blob
    return fetch


def _run(data_dir, title, blob, **kw):
    return ex.export(data_dir, sheet_id="sheet-id", fetch=_fetch(title, blob), today=TODAY, **kw)


class TestOnlyContentCounts:
    def test_unchanged_content_under_a_new_name_is_no_change(self, data_dir):
        r = _run(data_dir, NEW, _xlsx())
        assert r["changed"] is False
        assert sorted(p.name for p in data_dir.glob("*.xlsx")) == [f"{OLD}.xlsx"]
        assert (data_dir / "README.md").read_text(encoding="utf-8") == README

    def test_the_fingerprint_ignores_how_the_zip_was_packed(self, tmp_path):
        a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
        a.write_bytes(_xlsx())
        time.sleep(1.1)  # the zip and docProps carry timestamps to the second
        b.write_bytes(_xlsx())
        assert a.read_bytes() != b.read_bytes()
        assert ex.fingerprint(a) == ex.fingerprint(b)

    @pytest.mark.parametrize("change", [
        {"value": "Sol"},
        {"border": ORANGE},
        {"note": "Inferred: Pecaut & Mamajek"},
    ], ids=["value", "provenance-border", "note"])
    def test_a_value_border_or_note_is_a_change(self, data_dir, change):
        assert _run(data_dir, NEW, _xlsx(**change))["changed"] is True


class TestReplacingTheSnapshot:
    def test_new_file_named_after_the_workbook_and_old_one_gone(self, data_dir):
        r = _run(data_dir, NEW, _xlsx(value="Sol"))
        assert r["path"] == data_dir / f"{NEW}.xlsx"
        assert sorted(p.name for p in data_dir.glob("*.xlsx")) == [f"{NEW}.xlsx"]
        wb = openpyxl.load_workbook(r["path"], data_only=True)
        assert wb["Systems"]["A2"].value == "Sol"

    def test_readme_row_names_the_new_snapshot_and_nothing_else_moves(self, data_dir):
        _run(data_dir, NEW, _xlsx(value="Sol"))
        text = (data_dir / "README.md").read_text(encoding="utf-8")
        assert f"| `{NEW}.xlsx` | `{NEW}` | private (`sheet_id.txt`) | 2026-10-12 |" in text
        assert OLD not in text
        assert text.replace(ex._ROW.search(text).group(0), "") == \
            README.replace(ex._ROW.search(README).group(0), "")

    def test_force_replaces_identical_content(self, data_dir):
        r = _run(data_dir, NEW, _xlsx(), force=True)
        assert r["changed"] is True
        assert sorted(p.name for p in data_dir.glob("*.xlsx")) == [f"{NEW}.xlsx"]

    def test_force_under_the_same_name_keeps_the_file(self, data_dir):
        r = _run(data_dir, OLD, _xlsx(), force=True)
        assert r["changed"] is True
        assert sorted(p.name for p in data_dir.glob("*.xlsx")) == [f"{OLD}.xlsx"]

    @pytest.mark.parametrize("title", ["hy-brasil-20261010-2036", "../evil", "hybrasil-2026"])
    def test_a_name_off_the_versioning_rule_is_refused(self, data_dir, title):
        with pytest.raises(RuntimeError, match="hybrasil-YYYYMMDD-HHMM"):
            _run(data_dir, title, _xlsx(value="Sol"))
        assert sorted(p.name for p in data_dir.glob("*.xlsx")) == [f"{OLD}.xlsx"]


def test_main_reports_to_github_output(data_dir, tmp_path, monkeypatch):
    out = tmp_path / "gh_output"
    monkeypatch.setenv("GITHUB_OUTPUT", str(out))
    monkeypatch.setattr(ex, "export", lambda *a, **k: {"changed": True, "title": NEW, "path": "x"})
    assert ex.main(["--data-dir", str(data_dir)]) == 0
    assert out.read_text(encoding="utf-8") == f"changed=true\ntitle={NEW}\n"


def test_the_real_readme_has_the_row_this_script_rewrites():
    text = (ex.DATA_DIR / ex.README).read_text(encoding="utf-8")
    assert len(ex._ROW.findall(text)) == 1
